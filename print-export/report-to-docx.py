import sys
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
soup = BeautifulSoup(open(src).read(), 'html.parser')
doc = soup.select_one('.doc')
D = Document()
for s in D.sections:
    s.left_margin = s.right_margin = Inches(0.9)
    s.top_margin = s.bottom_margin = Inches(0.8)
st = D.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10.5); st.paragraph_format.space_after = Pt(6)
CH = RGBColor(0x28, 0x29, 0x24); MU = RGBColor(0x5a, 0x5a, 0x52)


def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexc); tcPr.append(sh)


def runs(p, node, bold=False, italic=False, size=None, color=None, font=None):
    for c in node.children:
        if isinstance(c, NavigableString):
            t = ' '.join(str(c).split('\n'))
            if not t: continue
            r = p.add_run(t); r.bold = bold or None; r.italic = italic or None
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
            if font: r.font.name = font
        elif c.name == 'br':
            p.add_run().add_break()
        else:
            runs(p, c, bold or c.name in ('b', 'strong'), italic or c.name in ('i', 'em'), size, color, font)


def para(node, **kw):
    p = D.add_paragraph(); runs(p, node, **kw); return p


for el in doc.children:
    if isinstance(el, NavigableString): continue
    cls = el.get('class') or []
    if 'mast' in cls:
        para(el.select_one('.lbl'), size=8.5, color=RGBColor(0x8a, 0x86, 0x65))
        para(el.find('h2'), bold=True, size=20, font='Georgia')
        para(el.select_one('.sub'), size=9.5, color=MU)
    elif el.name == 'p':
        if 'lede' in cls: para(el, size=11.5, font='Georgia')
        elif 'foot' in cls: para(el, size=8.5, color=MU)
        else: para(el)
    elif el.name == 'h3':
        p = para(el, bold=True, size=13, color=CH); p.paragraph_format.space_before = Pt(14)
    elif 'stats' in cls:
        stats = el.select('.stat'); t = D.add_table(rows=2, cols=len(stats)); t.style = 'Table Grid'
        for i, s in enumerate(stats):
            a = t.cell(0, i).paragraphs[0]; a.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = a.add_run(s.select_one('.n').get_text()); r.bold = True; r.font.size = Pt(18)
            b = t.cell(1, i).paragraphs[0]; b.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = b.add_run(s.select_one('.l').get_text(' ')); r.font.size = Pt(8.5); r.font.color.rgb = MU
            shade(t.cell(0, i), 'FCFBF9'); shade(t.cell(1, i), 'FCFBF9')
        D.add_paragraph()
    elif el.name == 'table':
        rows = el.find_all('tr'); n = len(rows[0].find_all(['th', 'td']))
        t = D.add_table(rows=len(rows), cols=n); t.style = 'Table Grid'
        for ri, tr in enumerate(rows):
            for ci, c in enumerate(tr.find_all(['th', 'td'])):
                cc = c.get('class') or []
                p = t.cell(ri, ci).paragraphs[0]
                runs(p, c, bold=(c.name == 'th' or 'pg' in cc), size=9.5)
                if c.name == 'th': shade(t.cell(ri, ci), 'EBE9E4')
                if 'n' in cc: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        D.add_paragraph()
    elif 'callout' in cls:
        t = D.add_table(rows=1, cols=1); c = t.cell(0, 0); shade(c, 'FFF6F6' if 'warn' in cls else 'EBE9E4')
        runs(c.paragraphs[0], el.select_one('.h'), bold=True, size=10)
        for pp in el.find_all('p'):
            runs(c.add_paragraph(), pp, size=10)
        D.add_paragraph()
    elif el.name == 'ul':
        for li in el.find_all('li'):
            runs(D.add_paragraph(style='List Bullet'), li)
D.save(out)
